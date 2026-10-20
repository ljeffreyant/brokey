import sys, math, warnings
warnings.simplefilter("ignore")
import openpyxl
a = openpyxl.load_workbook(sys.argv[1], data_only=True)   # my cached values
b = openpyxl.load_workbook(sys.argv[2], data_only=True)   # LibreOffice-calculated
f = openpyxl.load_workbook(sys.argv[1], data_only=False)
bad = 0; checked = 0
for ws in a.worksheets:
    wb_ = b[ws.title]; wf = f[ws.title]
    for row in wf.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('='):
                checked += 1
                va = ws[c.coordinate].value; vb = wb_[c.coordinate].value
                if va in (None, '') and vb in (None, ''): continue
                ok = False
                if isinstance(va, (int, float)) and isinstance(vb, (int, float)):
                    ok = math.isclose(va, vb, rel_tol=1e-9, abs_tol=1e-6)
                else:
                    ok = (str(va) == str(vb))
                    if not ok and hasattr(va, 'isoformat') and hasattr(vb, 'isoformat'): ok = va == vb
                if not ok:
                    bad += 1
                    if bad < 40: print('MISMATCH', ws.title, c.coordinate, c.value[:70], '| mine:', repr(va), '| LO:', repr(vb))
print('formulas checked', checked, 'mismatches', bad)
